#!/usr/bin/env python3
"""Render the weekly 'The Best AI is Yet to Come!' teaser slide for the AI Briefing.

Usage: python3 tools/make_slide.py spec.json out.pptx

spec.json:
{
  "week": "Week of 11 October 2026",
  "headline": "Cheaper models, smaller decisions",          (max 44 chars)
  "hooks": [                                                 (exactly 3)
    {"tag": "Delivery", "text": "Haiku 5.5 lists at a tenth of the old price"}   (tag max 16, text max 52)
  ],
  "cta": "Pick your role, press Play, listen on the way",   (max 52 chars)
  "url": "https://tonvlot1969.github.io/ai-briefing/ai-briefing.html"
}

One EPAM-branded content slide: Snow background, Night header bar, Calibri,
Cobalt and Slate accents, QR code to the app. Every text line fits on one line;
the script refuses text that is too long so the caller shortens the wording.
"""
import io
import json
import sys

import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

NIGHT = RGBColor(0x06, 0x06, 0x06)
SNOW = RGBColor(0xFB, 0xFA, 0xFA)
SEA = RGBColor(0x00, 0xF6, 0xFF)
COBALT = RGBColor(0x00, 0x47, 0xFF)
SLATE = RGBColor(0x4A, 0x71, 0xBD)
TINT = RGBColor(0xE9, 0xEF, 0xFF)

LIMITS = {"headline": 44, "tag": 16, "text": 52, "cta": 52, "week": 34}


def check(spec):
    errs = []
    for k in ("week", "headline", "cta"):
        if len(spec.get(k, "")) > LIMITS[k]:
            errs.append(f"{k} is {len(spec[k])} chars, max {LIMITS[k]}: {spec[k]!r}")
    hooks = spec.get("hooks", [])
    if len(hooks) != 3:
        errs.append(f"need exactly 3 hooks, got {len(hooks)}")
    for h in hooks:
        if len(h["tag"]) > LIMITS["tag"]:
            errs.append(f"tag too long ({len(h['tag'])}>{LIMITS['tag']}): {h['tag']!r}")
        if len(h["text"]) > LIMITS["text"]:
            errs.append(f"hook too long ({len(h['text'])}>{LIMITS['text']}): {h['text']!r}")
    text = json.dumps(spec, ensure_ascii=False)
    if "—" in text:
        errs.append("em-dash found, use a period or a comma")
    if errs:
        sys.exit("SLIDE SPEC REJECTED\n" + "\n".join(errs))


def box(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def main(spec_path, out_path):
    spec = json.load(open(spec_path, encoding="utf-8"))
    check(spec)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = SNOW

    # header bar
    rect(s, 0, 0, 13.333, 1.05, NIGHT)
    box(s, 0.6, 0.18, 8.4, 0.7, "AI Briefing  ·  " + spec["week"], 26, SNOW, bold=True)
    box(s, 8.9, 0.18, 3.85, 0.7, "The Best AI is Yet to Come!", 16, SEA, bold=True, align=PP_ALIGN.RIGHT)

    # headline
    box(s, 0.6, 1.45, 8.6, 0.8, spec["headline"], 34, COBALT, bold=True)

    # three hooks
    y = 2.6
    for h in spec["hooks"]:
        rect(s, 0.6, y, 8.6, 0.95, TINT, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.18
        rect(s, 0.6, y, 0.12, 0.95, COBALT)
        box(s, 0.95, y + 0.08, 8.0, 0.35, h["tag"].upper(), 14, SLATE, bold=True)
        box(s, 0.95, y + 0.4, 8.1, 0.48, h["text"], 21, NIGHT)
        y += 1.13

    # call to action
    box(s, 0.6, 6.15, 8.6, 0.6, spec["cta"], 20, NIGHT, bold=True)

    # QR card on the right
    rect(s, 9.75, 1.45, 3.0, 5.3, TINT, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.06
    qr = qrcode.QRCode(border=1, box_size=12, error_correction=qrcode.constants.ERROR_CORRECT_M)
    qr.add_data(spec["url"])
    qr.make(fit=True)
    img = qr.make_image(fill_color="#060606", back_color="#FBFAFA")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    s.shapes.add_picture(buf, Inches(9.95), Inches(1.7), Inches(2.6), Inches(2.6))
    box(s, 9.85, 4.5, 2.8, 0.5, "Scan and listen", 22, COBALT, bold=True, align=PP_ALIGN.CENTER)
    box(s, 9.85, 5.05, 2.8, 0.4, "NL or EN, 7 topics", 15, SLATE, align=PP_ALIGN.CENTER)
    box(s, 9.85, 5.5, 2.8, 0.4, "pick your role", 15, SLATE, align=PP_ALIGN.CENTER)
    box(s, 9.85, 6.1, 2.8, 0.4, "tonvlot1969.github.io/ai-briefing", 11, SLATE, align=PP_ALIGN.CENTER)

    prs.save(out_path)
    print("saved", out_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
